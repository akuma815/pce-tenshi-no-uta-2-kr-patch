#!/usr/bin/env python3
"""B트랙 B3 v2 — Switch판 최종 번역을 PCE 레코드에 run-정렬로 매칭.

핵심 사실 (2026-08-22 실측):
  - Switch 대사 행은 페이지(대화창) 단위이며 jp/ko 모두 master 계열 마크업을
    중괄호 없이 담고 있다: `#5フェイト：\\n#1本文`, `%F.` = {%F}+파라미터.
  - PCE master 레코드 = 여러 페이지를 `\\p` 로 이은 전체 메시지.
  - 페이지 분할·줄바꿈은 B4에서 한국어 기준으로 재구성하므로, B3의 책임은
    "레코드의 내용을 순서대로 완전하게 덮는 Switch ko 열람"이다.

매칭 알고리즘 (레코드별):
  R = 제어 제거·정규화한 레코드 전체 jp.
  1. R == 단일 행 norm → SWITCH_WHOLE
  2. R 을 접두사부터 덮는 연속 행 run(i..j) 탐색 (인접 행 norm 연결) → SWITCH_RUN
  3. 실패 시 구 번역(레코드 전체) 폴백 → OLD_FALLBACK
  4. 그것도 없으면 MISS (원문 유지)

출력 ko 는 master 표기({#x}, {%x}p, \\n, \\p)로 정규화해 기록한다.
"""

from __future__ import annotations

import csv
import hashlib
import json
import re
import unicodedata
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
MASTER = ROOT / "translation" / "full_extraction" / "master_text.csv"
TR = ROOT / "Translate_Original"
ASSIGN = ROOT / "translation" / "btrack_slot_assignment.json"

OUT_TSV = ROOT / "translation" / "btrack_b3_translation_map.tsv"
OUT_REVIEW = ROOT / "translation" / "btrack_b3_review.csv"
OUT_MANIFEST = ROOT / "translation" / "btrack_b3_manifest.json"

BSN = "\\" + "n"
BSP = "\\" + "p"

# master 표기 제어 토큰
MASTER_TOKEN_RE = re.compile(r"(\{#[0-9A-Fa-f]{1,2}\}|\{%[0-9A-Za-z]\}.|\\n|\\p)")
# Switch 셀 표기 제어 토큰 (중괄호 없음)
SWITCH_TOKEN_RE = re.compile(r"(#[0-9A-Fa-f]|%[0-9A-Za-z].|\\n)")

STRIP_RE = re.compile(r"[\s、。！？!?…・．，「」『』（）()：:～〜ー―]")


def norm_text(s: str) -> str:
    return STRIP_RE.sub("", unicodedata.normalize("NFKC", str(s or "")))


def master_norm(s: str) -> str:
    """master 표기 문자열에서 제어 제거 후 정규화."""
    return norm_text(MASTER_TOKEN_RE.sub("", s or ""))


def switch_norm(s: str) -> str:
    """Switch 셀 표기 문자열에서 제어 제거 후 정규화."""
    return norm_text(SWITCH_TOKEN_RE.sub("", str(s or "")))


def switch_to_master_notation(s: str) -> str:
    """#5 → {#5}, %F. → {%F}., \\n 은 그대로."""
    s = str(s or "")
    s = re.sub(r"%([0-9A-Za-z])(.)", r"{%\1}\2", s)
    s = re.sub(r"#([0-9A-Fa-f])", r"{#\1}", s)   # 제어는 항상 # + 한 자리 hex
    return s


# ---- Switch 행 적재 (순서 보존) ----------------------------------------------

def load_switch_rows():
    rows = []          # (norm, ko_master, source, row_no)
    counts = {}

    def add(jp, ko, src, no):
        n = switch_norm(jp)
        ko_m = switch_to_master_notation(ko) if ko is not None else None
        rows.append((n, ko_m, src, no))

    for xlsx in ("천사의시2_대사_번역.xlsx", "천사의시2_크레딧_번역.xlsx"):
        wb = load_workbook(TR / xlsx, read_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        hd = [str(c) for c in next(it)]
        ji, ki = hd.index("jp"), hd.index("ko")
        n0 = len(rows)
        for no, r in enumerate(it):
            if r and r[ji] is not None:
                add(r[ji], r[ki], xlsx, no)
        counts[xlsx] = len(rows) - n0
        wb.close()

    with (TR / "visualscene2_paired.csv").open(encoding="utf-8-sig", newline="") as f:
        n0 = len(rows)
        for no, row in enumerate(csv.DictReader(f)):
            if row.get("jp"):
                add(row["jp"], row.get("ko"), "visualscene2_paired.csv", no)
        counts["visualscene2_paired.csv"] = len(rows) - n0

    # 색인: 전체 일치 / 접두사 후보(앞 8자)
    whole = {}
    pre8 = {}
    for idx, (n, ko, src, no) in enumerate(rows):
        if not n or ko is None:
            continue
        whole.setdefault(n, idx)
        pre8.setdefault(n[:8], []).append(idx)
    return rows, whole, pre8, counts


# ---- run 정렬 ----------------------------------------------------------------

def align_record(R: str, rows, whole, pre8, max_run=24):
    """R 전체를 덮는 (행 인덱스 리스트) 반환. 실패 시 None."""
    if R in whole:
        return [whole[R]]
    best = None
    for i in pre8.get(R[:8], []):
        n0 = rows[i][0]
        if not R.startswith(n0):
            continue
        acc = n0
        run = [i]
        j = i
        while len(acc) < len(R) and len(run) < max_run and j + 1 < len(rows):
            j += 1
            nj, koj = rows[j][0], rows[j][1]
            if not nj or koj is None:
                continue                      # 빈 행 건너뜀
            if R.startswith(acc + nj):
                acc += nj
                run.append(j)
            else:
                break
        if acc == R:
            if best is None or len(run) < len(best):
                best = run
    return best


PAGE_SPK_RE = re.compile(
    r"^((?:\{%[0-9A-Za-z]\}.|\{#[0-9A-Fa-f]{1,2}\})*)"      # prefix controls
    r"([^\\{}]{1,12}[：:])\\n\{#1\}")            # SPEAKER：BSN{#1}


LEAD_CTRL_RE = re.compile(r"^(?:\{%[0-9A-Za-z]\}.|\{#[0-9A-Fa-f]{1,2}\})+")


def parse_page(page: str):
    """(prefix, speaker_with_colon|None, body)"""
    m = PAGE_SPK_RE.match(page)
    if m:
        return m.group(1), m.group(2), page[m.end():]
    m2 = re.match(r"^((?:\{%[0-9A-Za-z]\}.|\{#[0-9A-Fa-f]{1,2}\})*)", page)
    return m2.group(1), None, page[m2.end():]


def strip_row_markup(ko: str):
    """행 ko 가 자체 화자 마크업을 가지면 (이름, 본문) 분리."""
    m = PAGE_SPK_RE.match(ko)
    if m:
        return m.group(2).rstrip("：:"), ko[m.end():]
    return None, ko


def assemble_record(original: str, run_rows):
    """원본 페이지·화자 골격 위에 Switch 행 ko 를 재조립. 실패 시 (None, False)."""
    queue = list(run_rows)
    out_pages = []
    for page in original.split(BSP):
        prefix, speaker_jp, body_jp = parse_page(page)
        body_norm = master_norm(body_jp)
        spk_norm = norm_text(speaker_jp.rstrip("：:")) if speaker_jp else ""
        speaker_ko = None
        body_parts = []
        if speaker_jp and queue and queue[0][0] == spk_norm:
            speaker_ko = strip_row_markup(queue.pop(0)[1])[1]
            speaker_ko = speaker_ko.replace("：", "").replace(":", "").strip()
        if body_norm == "":
            body_ko = body_jp                # 구두점 전용 본문(・・・)은 원문 통과
        else:
            acc = ""
            while acc != body_norm:
                if not queue:
                    return None, False
                n_, ko_ = queue.pop(0)
                nm, bd = strip_row_markup(ko_)
                if acc == "" and speaker_jp and speaker_ko is None and n_ == spk_norm + body_norm:
                    if nm:
                        speaker_ko = nm      # 행이 화자+본문 통짜 (자체 마크업)
                    else:
                        bd = LEAD_CTRL_RE.sub("", bd, count=1)
                    body_parts.append(bd)
                    acc = body_norm
                    break
                if not body_norm.startswith(acc + n_):
                    return None, False
                if acc == "":
                    bd = LEAD_CTRL_RE.sub("", bd, count=1)   # 선두 제어는 prefix 가 공급
                acc += n_
                if nm and speaker_ko is None:
                    speaker_ko = nm
                body_parts.append(bd)
            body_ko = BSN.join(x for x in body_parts if x != "")
        if speaker_jp:
            if speaker_ko is None:
                speaker_ko = speaker_jp.rstrip("：:")
            out_pages.append(prefix + speaker_ko + "：" + BSN + "{#1}" + body_ko)
        else:
            out_pages.append(prefix + body_ko)
    if queue:
        return None, False
    return BSP.join(out_pages), True


def main() -> None:
    rows_sw, whole, pre8, counts = load_switch_rows()
    master_rows = list(csv.DictReader(MASTER.open(encoding="utf-8-sig", newline="")))

    results = {}
    agg = Counter()
    review = []
    used_rows = set()
    jp_re = re.compile(r"[぀-ヿ一-鿿]")

    for r in master_rows:
        if r["duplicate_of"]:
            continue
        original = r["original"]
        if not original or not jp_re.search(original):
            continue
        R = master_norm(original)
        run = align_record(R, rows_sw, whole, pre8) if R else None
        if run:
            run_rows = [(rows_sw[i][0], rows_sw[i][1]) for i in run]
            ko, ok = assemble_record(original, run_rows)
            if not ok:
                ko = BSP.join(rows_sw[i][1] for i in run)
            used_rows.update(run)
            status = ("SWITCH_WHOLE" if len(run) == 1 else "SWITCH_RUN")                 + ("" if ok else "_JOIN")
        elif r.get("translation"):
            ko = r["translation"]
            status = "OLD_FALLBACK"
        else:
            ko = original
            status = "MISS"
        results[r["id"]] = {
            "id": r["id"], "category": r["category"], "status": status,
            "run_len": len(run) if run else 0, "ko": ko,
        }
        agg[status] += 1
        if status not in ("SWITCH_WHOLE", "SWITCH_RUN"):
            review.append((r["id"], r["category"], status,
                           original[:150], ko[:150]))

    inherited = 0
    for r in master_rows:
        d = r["duplicate_of"]
        if d and d in results:
            results[r["id"]] = {**results[d], "id": r["id"], "inherits": d}
            inherited += 1

    with OUT_TSV.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, delimiter="\t")
        w.writerow(["id", "category", "status", "inherits", "run_len", "ko"])
        for rid in sorted(results):
            e = results[rid]
            w.writerow([e["id"], e.get("category", ""), e["status"],
                        e.get("inherits", ""), e.get("run_len", ""), e["ko"]])
    with OUT_REVIEW.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["id", "category", "status", "original", "ko"])
        w.writerows(review)

    assigned = set(json.loads(ASSIGN.read_text(encoding="utf-8"))["assignment"])
    missing_glyphs = Counter()
    for e in results.values():
        if "inherits" in e:
            continue
        for ch in e["ko"]:
            if 0xAC00 <= ord(ch) <= 0xD7A3 and ch not in assigned:
                missing_glyphs[ch] += 1

    manifest = {
        "format": "BTRACK-B3-TRANSLATION-MAP",
        "version": 2,
        "switch_sources": counts,
        "records_processed": len([1 for e in results.values() if "inherits" not in e]),
        "records_inherited": inherited,
        "record_status": dict(agg),
        "switch_rows_used": len(used_rows),
        "switch_rows_total": len(rows_sw),
        "review_rows": len(review),
        "missing_glyphs": {c: n for c, n in missing_glyphs.most_common()},
        "tsv_sha256": hashlib.sha256(OUT_TSV.read_bytes()).hexdigest(),
    }
    OUT_MANIFEST.write_text(json.dumps(manifest, ensure_ascii=False, indent=1),
                            encoding="utf-8")
    print(json.dumps({k: manifest[k] for k in
                      ("records_processed", "records_inherited", "record_status",
                       "switch_rows_used", "switch_rows_total", "review_rows")},
                     ensure_ascii=False, indent=1))
    if missing_glyphs:
        print("경고: 미배정 글리프", len(missing_glyphs), "종:",
              "".join(list(missing_glyphs)[:60]))


if __name__ == "__main__":
    main()
