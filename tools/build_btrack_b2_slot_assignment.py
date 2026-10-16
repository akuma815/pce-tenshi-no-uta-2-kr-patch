#!/usr/bin/env python3
"""B트랙 B2 — 전체 한글 슬롯 배정 + 풀 시스템카드 빌드.

입력:
  - Switch판 최종 번역 (Translate_Original/: 대사 xlsx, 크레딧 xlsx, paired csv 2종)
  - Galmuri11-Condensed 9x12 비트맵 시트 (가~힣 유니코드순, 64열, 획=비검정)
  - 원본 시스템카드 (mode-1 폰트: kanji 슬롯 453..3417 = level-1 한자 2,965)
  - master_text.csv 원문 코퍼스 (게임이 실제 쓰는 한자 → 미사용 한자 슬롯 우선 배정)

배정 규칙 (fail-closed):
  1. 카드 UI 한자(백업 관리 화면 실측: 削除全体初期化終了)는 절대 배정하지 않는다.
  2. 게임 원문에 등장하지 않는 한자 슬롯부터 배정한다 (개발 중 혼재 화면 최소화).
  3. 부족하면 게임 원문 등장 빈도가 낮은 한자 슬롯 순으로 배정한다.
  4. 한글 필요 글리프 수가 가용 슬롯을 넘으면 빌드 실패.

출력:
  - translation/btrack_slot_assignment.json  (한글 ↔ SJIS/slot 전체 매핑)
  - probes/SystemCard3_btrack_full.pce       (풀 한글 카드, 헤더 없는 262,144B)
  - translation/btrack_b2_manifest.json
"""

from __future__ import annotations

import csv
import hashlib
import json
import sys
from collections import Counter
from pathlib import Path

from PIL import Image
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))

BIOS_SRC = ROOT / "bios" / "Super CD-ROM2 System Ver. 3.00 (J).pce"
EXPECTED_BIOS_SHA1 = "79f5ff55dd10187c7fd7b8daab0b3ffbd1f56a2c"
FONT_PNG = Path(r"G:\Downloads\Korean Traslate Project\Font\Galmuri11-Condensed_9x12.png")
TR = ROOT / "Translate_Original"
MASTER = ROOT / "translation" / "full_extraction" / "master_text.csv"

OUT_ASSIGN = ROOT / "translation" / "btrack_slot_assignment.json"
OUT_CARD = ROOT / "probes" / "SystemCard3_btrack_full.pce"
OUT_MANIFEST = ROOT / "translation" / "btrack_b2_manifest.json"

FONT_BASE = 0x30000
GLYPH_BYTES = 18
KANJI_SLOT_MIN, KANJI_SLOT_MAX = 453, 3417   # 亜..腕 level-1 실측 구간
CARD_UI_KANJI = "削除全体初期化終了"          # 백업 관리 화면 실측 (확장 가능)
GLYPH_X_OFFSET = 1                            # 9px 글리프를 12px 셀에 배치


def sha256_bytes(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()


# ---- Switch 최종본에서 한국어 코퍼스 수집 -----------------------------------

def collect_korean() -> tuple[Counter, Counter, dict]:
    hangul = Counter()
    other = Counter()          # 한글·ASCII 외 문자 (구두점 등 — 보고용)
    sources = {}

    def feed(text, src):
        if not text:
            return
        for ch in str(text):
            o = ord(ch)
            if 0xAC00 <= o <= 0xD7A3:
                hangul[ch] += 1
            elif o >= 0x80 and ch not in "\n\r\t":
                other[ch] += 1

    for xlsx, col in (("천사의시2_대사_번역.xlsx", "ko"), ("천사의시2_크레딧_번역.xlsx", "ko")):
        wb = load_workbook(TR / xlsx, read_only=True)
        ws = wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header = [str(c) for c in next(rows)]
        ki = header.index(col)
        n = 0
        for r in rows:
            if r and r[ki]:
                feed(r[ki], xlsx)
                n += 1
        wb.close()
        sources[xlsx] = n
    for name in ("visualscene2_paired.csv", "outgameinfo2_paired.csv"):
        with (TR / name).open(encoding="utf-8-sig", newline="") as f:
            n = 0
            for row in csv.DictReader(f):
                feed(row.get("ko"), name)
                n += 1
            sources[name] = n
    # B3 재조립 결과(구 번역 폴백 포함) — 존재하면 코퍼스에 합산
    b3 = ROOT / "translation" / "btrack_b3_translation_map.tsv"
    if b3.exists():
        with b3.open(encoding="utf-8", newline="") as f:
            n = 0
            for row in csv.DictReader(f, delimiter="	"):
                if not row.get("inherits"):
                    feed(row.get("ko"), "btrack_b3")
                    n += 1
            sources["btrack_b3_translation_map.tsv"] = n
    return hangul, other, sources


# ---- 게임 원문 한자 사용 통계 ------------------------------------------------

def game_kanji_usage() -> Counter:
    used = Counter()
    with MASTER.open(encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            for ch in row.get("original") or "":
                if 0x4E00 <= ord(ch) <= 0x9FFF:
                    used[ch] += 1
    return used


# ---- 폰트 시트 추출 ----------------------------------------------------------

def load_font_sheet():
    im = Image.open(FONT_PNG).convert("RGB")
    if im.size != (576, 2100):
        raise ValueError(f"예상 밖 폰트 시트 크기: {im.size}")
    return im


def render_hangul(im, ch: str) -> bytes:
    idx = ord(ch) - 0xAC00
    cx, cy = (idx % 64) * 9, (idx // 64) * 12
    rows = []
    for y in range(12):
        row = 0
        for x in range(9):
            if im.getpixel((cx + x, cy + y)) != (0, 0, 0):
                row |= 1 << (11 - (x + GLYPH_X_OFFSET))
        rows.append(row)
    if not any(rows):
        raise ValueError(f"폰트 시트에 빈 글리프: {ch} (idx {idx})")
    out = bytearray()
    for y in range(0, 12, 2):
        r0, r1 = rows[y], rows[y + 1]
        out.extend((r0 >> 4, ((r0 & 0x0F) << 4) | (r1 >> 8), r1 & 0xFF))
    return bytes(out)


def main() -> None:
    bios_raw = BIOS_SRC.read_bytes()
    bios = bytearray(bios_raw[512:])
    if hashlib.sha1(bytes(bios)).hexdigest() != EXPECTED_BIOS_SHA1:
        raise ValueError("시스템카드 SHA-1 불일치")

    from analyze_system_card_font_slots import glyph_index  # noqa: E402
    if glyph_index("出".encode("shift_jis"), bytes(bios)) != 1722:
        raise AssertionError("glyph_index 회귀")

    # ---- kanji 슬롯 전수 목록 (slot -> 대표 SJIS/한자) ----------------------
    slot_info = {}
    for lead in list(range(0x88, 0xA0)) + list(range(0xE0, 0xF0)):
        for trail in list(range(0x40, 0x7F)) + list(range(0x80, 0xFD)):
            code = bytes((lead, trail))
            idx = glyph_index(code, bytes(bios))
            if idx is None or not (KANJI_SLOT_MIN <= idx <= KANJI_SLOT_MAX):
                continue
            if idx in slot_info:
                continue
            try:
                kanji = code.decode("shift_jis")
            except UnicodeDecodeError:
                continue
            slot = bios[FONT_BASE + idx * GLYPH_BYTES:FONT_BASE + (idx + 1) * GLYPH_BYTES]
            if all(b in (0x00, 0xFF) for b in slot):
                continue                      # 실존 글리프만
            slot_info[idx] = (code, kanji)
    total_slots = len(slot_info)

    # ---- 제외: 카드 UI 한자 -------------------------------------------------
    ui_excluded = {}
    for ch in CARD_UI_KANJI:
        idx = glyph_index(ch.encode("shift_jis"), bytes(bios))
        if idx is None:
            raise AssertionError(f"카드 UI 한자 인덱스 실패: {ch}")
        ui_excluded[idx] = ch
        slot_info.pop(idx, None)

    # ---- 수요·공급 ----------------------------------------------------------
    hangul, other_chars, sources = collect_korean()
    need = [ch for ch, _ in hangul.most_common()]
    if len(need) > len(slot_info):
        raise SystemExit(f"FAIL: 한글 {len(need)}자 > 가용 슬롯 {len(slot_info)}")

    usage = game_kanji_usage()
    # 게임 미사용 한자 슬롯 먼저, 그다음 저빈도 순 — 각 그룹 안은 slot 오름차순
    order = sorted(slot_info.items(), key=lambda kv: (usage.get(kv[1][1], 0), kv[0]))

    # ---- 배정 + 카드 빌드 ---------------------------------------------------
    im = load_font_sheet()
    assignment = {}
    card_slots_unused_game = 0
    for ch, (idx, (code, kanji)) in zip(need, order):
        glyph = render_hangul(im, ch)
        off = FONT_BASE + idx * GLYPH_BYTES
        bios[off:off + GLYPH_BYTES] = glyph
        if usage.get(kanji, 0) == 0:
            card_slots_unused_game += 1
        assignment[ch] = {
            "sjis": code.hex().upper(), "glyph_index": idx,
            "original_kanji": kanji, "kanji_game_freq": usage.get(kanji, 0),
            "ko_freq": hangul[ch],
        }
    OUT_CARD.write_bytes(bytes(bios))

    OUT_ASSIGN.write_text(json.dumps({
        "format": "BTRACK-SLOT-ASSIGNMENT",
        "version": 1,
        "font_sheet": str(FONT_PNG),
        "glyph_x_offset": GLYPH_X_OFFSET,
        "card": str(OUT_CARD),
        "assignment": assignment,
    }, ensure_ascii=False, indent=1), encoding="utf-8")

    manifest = {
        "format": "BTRACK-B2-FULL-CARD",
        "runtime_status": "PENDING",
        "input": {
            "original_bios_sha1": EXPECTED_BIOS_SHA1,
            "font_png_sha256": sha256_bytes(FONT_PNG.read_bytes()),
            "korean_sources": sources,
        },
        "supply": {
            "kanji_slots_total": total_slots,
            "card_ui_excluded": {str(k): v for k, v in ui_excluded.items()},
        },
        "demand": {
            "unique_hangul": len(need),
            "hangul_occurrences": sum(hangul.values()),
            "non_hangul_chars": "".join(sorted(other_chars,
                                               key=lambda c: -other_chars[c]))[:120],
        },
        "assignment_stats": {
            "assigned": len(assignment),
            "into_game_unused_kanji_slots": card_slots_unused_game,
            "into_game_used_kanji_slots": len(assignment) - card_slots_unused_game,
            "slots_remaining": len(slot_info) - len(assignment),
        },
        "card_sha256": sha256_bytes(bytes(bios)),
        "runtime_gate": [
            "부트·백업 관리 화면 원본과 동일 (자동)",
            "원본 게임 전체 회귀 — 일본어 텍스트 중 배정된 한자만 한글로 보임 (혼재 허용)",
        ],
    }
    OUT_MANIFEST.write_text(json.dumps(manifest, ensure_ascii=False, indent=1),
                            encoding="utf-8")
    print(json.dumps({
        "status": "BUILD_OK",
        "unique_hangul": len(need),
        "kanji_slots_available": len(slot_info),
        "into_game_unused": card_slots_unused_game,
        "into_game_used": len(assignment) - card_slots_unused_game,
        "slots_remaining": len(slot_info) - len(assignment),
        "card": OUT_CARD.name,
        "card_sha256": manifest["card_sha256"][:16],
    }, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
